"""Builds the poster numbers and charts from the case sheet.
Run:  python make_outputs.py   ->  outputs/Poster_Results.xlsx + outputs/*.png
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

import engine as e

OUT = Path(__file__).parent / "outputs"
OUT.mkdir(exist_ok=True)
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.spines.top": False,
                     "axes.spines.right": False})
BLUE, GREY, RED = "#2E6DA4", "#B0B7C3", "#C0504D"

df = e.load_cases()
table, model, scores = e.train(df)
df["Score"] = scores
df["Probability"] = [e.probability(model, s) for s in scores]
sar = df[df.y == 1]
n_s = len(sar)

freq = pd.DataFrame({
    "Symptom / finding": e.CONFIG["features"],
    "Sarcoidosis patients with it": [(sar[f] == "Yes").sum() for f in e.CONFIG["features"]],
})
freq["Out of"] = n_s
freq["%"] = freq["Sarcoidosis patients with it"] / n_s
freq = freq.sort_values("Sarcoidosis patients with it", ascending=False).reset_index(drop=True)
freq.insert(0, "Rank", range(1, len(freq) + 1))

pts = table[["Finding", "Points", "% in sarcoidosis", "% in other diseases",
             "Starting points (clinical knowledge)"]].copy()
pts = pts.sort_values("Points", ascending=False).reset_index(drop=True)
pts.insert(0, "Rank", range(1, len(pts) + 1))

labs = pd.DataFrame([{"Lab": k, "Unit": v["unit"], "Normal low": v["low"], "Normal high": v["high"],
                      "Counts when": "above normal" if v["direction"] == "high" else "below normal",
                      "Max points": v["max_points"]} for k, v in e.CONFIG["labs"].items()])
tiers = pd.DataFrame({"How far outside normal": ["Up to 1.5x the limit", "1.5x to 2x", "More than 2x"],
                      "Share of the lab's max points": ["1/3", "2/3", "Full"]})

age = sar["Age (years)"].dropna()
dt = sar["Diagnostic time (months)"].dropna()
stats = pd.DataFrame([
    ("Total cases", len(df)),
    ("Sarcoidosis cases", n_s),
    ("Other diseases (look-alikes)", len(df) - n_s),
    ("Sarcoidosis – average age (years)", round(age.mean(), 1)),
    ("Sarcoidosis – age range", f"{age.min():g} – {age.max():g}"),
    ("Sarcoidosis – most common age group", "40–59"),
    ("Sarcoidosis – female", f"{(sar.Gender == 'Female').sum()} of {n_s}"),
    ("Sarcoidosis – male", f"{(sar.Gender == 'Male').sum()} of {n_s}"),
    ("Sarcoidosis – average time to diagnosis (months)", round(dt.mean(), 1)),
    ("Sarcoidosis – median time to diagnosis (months)", round(dt.median(), 1)),
    ("Sarcoidosis – longest time to diagnosis (months)", dt.max()),
    ("Look-alike diseases in the data", ", ".join(sorted({d.strip().lower(): d.strip() for d in df[df.y == 0]["Final diagnosis"]}.values(), key=str.lower))),
], columns=["Item", "Value"])
bins = pd.cut(age, [0, 18, 40, 60, 120], right=False, labels=["Under 18", "18–39", "40–59", "60+"])
top_group = bins.value_counts().idxmax()
stats.loc[stats.Item == "Sarcoidosis – most common age group", "Value"] = top_group

cases = df[["Case #", "Final diagnosis", "Sarcoidosis?", "Score", "Probability"]].copy()

xlsx = OUT / "Poster_Results.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
    stats.to_excel(w, sheet_name="Summary", index=False)
    freq.to_excel(w, sheet_name="Most frequent symptoms", index=False)
    pts.to_excel(w, sheet_name="Points per finding", index=False)
    labs.to_excel(w, sheet_name="Lab points", index=False)
    tiers.to_excel(w, sheet_name="Lab points", index=False, startrow=len(labs) + 3)
    cases.to_excel(w, sheet_name="Score per case", index=False)

wb = load_workbook(xlsx)
for ws in wb:
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E6DA4")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font is None or not c.font.bold:
                c.font = Font(name="Arial")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
for name, cols in (("Most frequent symptoms", ["E"]), ("Points per finding", ["D", "E"]),
                   ("Score per case", ["E"])):
    ws = wb[name]
    for col in cols:
        for c in ws[col][1:]:
            c.number_format = "0%"
lp = wb["Lab points"]
r = len(labs) + 4
for c in lp[r]:
    if c.value:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E6DA4")
lp.cell(len(labs) + 3, 1, "Lab severity rule").font = Font(name="Arial", bold=True)
wb.save(xlsx)

top = freq.head(10).iloc[::-1]
fig, ax = plt.subplots(figsize=(8, 5))
ax.barh(top["Symptom / finding"], top["Sarcoidosis patients with it"], color=BLUE)
for i, v in enumerate(top["Sarcoidosis patients with it"]):
    ax.text(v + 0.15, i, f"{v} of {n_s}", va="center", fontsize=10)
ax.set_xlabel(f"Number of sarcoidosis patients (out of {n_s})")
ax.set_title("Most frequent findings in sarcoidosis", loc="left", fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "1_most_frequent_symptoms.png", dpi=200)

p = pts.iloc[::-1]
fig, ax = plt.subplots(figsize=(8, 7))
ax.barh(p["Finding"], p["Points"], color=[BLUE if v > 0 else (RED if v < 0 else GREY) for v in p["Points"]])
ax.axvline(0, color="#555", lw=0.8)
ax.set_xlabel("Points (negative = points to a look-alike disease)")
ax.set_title("Points given to each finding", loc="left", fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "2_points_per_finding.png", dpi=200)

fig, ax = plt.subplots(figsize=(6, 4))
counts = bins.value_counts().reindex(["Under 18", "18–39", "40–59", "60+"])
ax.bar(counts.index, counts.values, color=BLUE)
for i, v in enumerate(counts.values):
    ax.text(i, v + 0.1, str(v), ha="center")
ax.set_ylabel("Patients")
ax.set_title("Age of sarcoidosis patients", loc="left", fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "3_age_groups.png", dpi=200)

fig, ax = plt.subplots(figsize=(7, 4))
d = sar.dropna(subset=["Diagnostic time (months)"]).sort_values("Diagnostic time (months)")
ax.bar(range(len(d)), d["Diagnostic time (months)"], color=BLUE)
ax.axhline(dt.median(), color=RED, ls="--", lw=1.2)
ax.text(-0.4, dt.median() + 4, f"median {dt.median():g} months", color=RED, ha="left")
ax.set_xticks([])
ax.set_xlabel("Each bar = one sarcoidosis patient")
ax.set_ylabel("Months until diagnosis")
ax.set_title("How long sarcoidosis took to diagnose", loc="left", fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "4_time_to_diagnosis.png", dpi=200)
print("saved", xlsx)
